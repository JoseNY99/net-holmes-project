import io
import json
import requests
from io import BytesIO
from PIL import Image as PILImage
from reportlab.platypus import Image as RLImage
from docx import Document
from docx.shared import Pt, Inches
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.enums import TA_CENTER
from sqlalchemy.orm import Session
from sqlalchemy import or_

from reportlab.platypus import Table, TableStyle
from reportlab.lib import colors


def generate_cbn_investigation_pdf(db, reporte, perfil):
    """Lazy import to avoid circular dependency."""
    from services.cbn_pdf import generate_cbn_investigation_pdf as _gen
    return _gen(db, reporte, perfil)

def fetch_and_compress_image(reporte):
    try:
        from pathlib import Path
        base_dir = Path(__file__).resolve().parent.parent
        local_evidencias = base_dir / "static" / "evidencias" / f"reporte_{reporte.id}.jpg"
        local_screenshots = base_dir / "static" / "screenshots" / f"reporte_{reporte.id}.jpg"
        
        target_file = None
        if local_evidencias.exists() and local_evidencias.stat().st_size > 1000:
            target_file = local_evidencias
        elif local_screenshots.exists() and local_screenshots.stat().st_size > 1000:
            target_file = local_screenshots
            
        if target_file:
            img = PILImage.open(target_file)
            if img.mode in ("RGBA", "P"):
                img = img.convert("RGB")
            max_size = (800, 800)
            img.thumbnail(max_size, PILImage.Resampling.LANCZOS)
            compressed_io = BytesIO()
            img.save(compressed_io, format="JPEG", quality=50, optimize=True)
            compressed_io.seek(0)
            return compressed_io

        plat = str(reporte.plataforma).lower()
        plat_data = None
        if plat == "facebook" and hasattr(reporte, "facebook_data"):
            plat_data = reporte.facebook_data
        elif plat == "instagram" and hasattr(reporte, "instagram_data"):
            plat_data = reporte.instagram_data
        elif plat == "tiktok" and hasattr(reporte, "tiktok_data"):
            plat_data = reporte.tiktok_data
        elif plat == "mercadolibre" and hasattr(reporte, "mercadolibre_data"):
            plat_data = reporte.mercadolibre_data
            
        image_url = None
        if plat_data and hasattr(plat_data, "image_urls") and plat_data.image_urls:
            try:
                urls = json.loads(plat_data.image_urls)
                if isinstance(urls, list) and len(urls) > 0:
                    image_url = urls[0]
            except:
                if isinstance(plat_data.image_urls, str) and plat_data.image_urls.startswith("http"):
                    image_url = plat_data.image_urls
                    
        # Fallback to general image_urls if exist in reporte
        if not image_url and hasattr(reporte, "image_urls") and reporte.image_urls:
            try:
                urls = json.loads(reporte.image_urls)
                if isinstance(urls, list) and len(urls) > 0:
                    image_url = urls[0]
            except:
                pass
                
        if not image_url:
            return None
            
        # Download
        response = requests.get(image_url, timeout=5)
        if response.status_code != 200:
            return None
            
        # Compress using PIL
        img = PILImage.open(BytesIO(response.content))
        if img.mode in ("RGBA", "P"):
            img = img.convert("RGB")
            
        max_size = (800, 800)
        img.thumbnail(max_size, PILImage.Resampling.LANCZOS)
        
        compressed_io = BytesIO()
        img.save(compressed_io, format="JPEG", quality=40, optimize=True)
        compressed_io.seek(0)
        
        return compressed_io
    except Exception as e:
        print(f"Error procesando imagen para export: {e}")
        return None

def generate_investigation_pdf(db: Session, reporte, perfil):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    
    title_style = styles['Heading1']
    title_style.alignment = TA_CENTER
    normal_style = styles['Normal']
    
    Story = []
    
    Story.append(Paragraph(f"INTELLIGENCE REPORT — {reporte.pais.upper()}", title_style))
    Story.append(Spacer(1, 12))
    Story.append(Paragraph(f"<b>Ref:</b> {reporte.referencia_interna} | <b>ID producto:</b> {reporte.id}", normal_style))
    Story.append(Spacer(1, 12))
    
    Story.append(Paragraph("<b>COMPANY OVERVIEW</b>", styles['Heading2']))
    Story.append(Spacer(1, 6))
    
    username = reporte.titular_cuenta
    arroba = ""
    username_clean = username.replace("@", "") if username else "Desconocido"
    owner_display = f"@{username_clean}"
    
    data_fase1 = [
        ["Campo (Fase 1: Recolección Automática)", "Valor detectado"],
        ["País", reporte.pais.upper()],
        ["Referencia interna", reporte.referencia_interna],
        ["Owner (Titular)", owner_display],
        ["Social Networks / Plataforma", reporte.plataforma],
        ["Website / URL evidencia", Paragraph(f"<a href='{reporte.url_evidencia}' color='blue'>{reporte.url_evidencia}</a>", normal_style)],
        ["Fecha de escaneo", reporte.fecha_captura.strftime("%Y-%m-%d %H:%M:%S") if reporte.fecha_captura else "N/D"],
        ["Proposed fees (Precio)", f"USD {reporte.precio_publico}" if reporte.precio_publico else "USD —"],
        ["Catalogue (Producto)", reporte.producto_catalogo],
        ["Marca afectada", reporte.marca_afectada],
        ["Sales (Hallazgos)", Paragraph(f"Según la investigación, se detectó oferta de productos con indicio de infracción a la marca {reporte.marca_afectada}. Origen: Automático.", normal_style)],
        ["Type of Seller", "Por determinar"],
        ["Estimated quantity", "Por determinar"],
        ["Address Evidence", reporte.pais]
    ]

    data_fase2 = [
        ["Campo (Fase 2: Relevamiento de Identidad)", "Valor detectado"],
        ["DNI / Cédula", getattr(perfil, 'dni_cedula', None) or "—"],
        ["CUIT / CUIL / RUC", getattr(perfil, 'cuit_cuil_ruc', None) or "—"],
        ["Banco", getattr(perfil, 'banco', None) or "—"],
        ["Alias bancario / MercadoPago", getattr(perfil, 'alias_bancario', None) or "—"],
        ["CVU / CBU", getattr(perfil, 'cvu_cbu', None) or "—"],
        ["Address (Dirección exacta)", getattr(perfil, 'direccion_exacta', None) or "—"],
        ["Telephone Number (Privado)", getattr(perfil, 'telefono_privado', None) or "—"],
        ["Email", getattr(perfil, 'email', None) or "—"],
        ["Precio final cotizado", getattr(perfil, 'precio_final_cotizado', None) or "—"]
    ]

    def _create_table(data_matrix):
        t = Table(data_matrix, colWidths=[200, 300])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#1A202C")),
            ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTNAME', (0,1), (0,-1), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,-1), 9),
            ('BOTTOMPADDING', (0,0), (-1,-1), 8),
            ('TOPPADDING', (0,0), (-1,-1), 8),
            ('GRID', (0,0), (-1,-1), 0.5, colors.lightgrey),
            ('VALIGN', (0,0), (-1,-1), 'TOP')
        ]))
        return t

    Story.append(Paragraph("<b>FASE 1 DE LA INVESTIGACIÓN (RECOLECCIÓN)</b>", styles['Heading3']))
    Story.append(Spacer(1, 6))
    Story.append(_create_table(data_fase1))
    
    Story.append(Spacer(1, 20))
    Story.append(Paragraph("<b>FASE 2 DE RELEVAMIENTO (IDENTIDAD)</b>", styles['Heading3']))
    Story.append(Spacer(1, 6))
    Story.append(_create_table(data_fase2))
    
    Story.append(Spacer(1, 20))
    Story.append(Paragraph("<b>ADDRESS / PHOTO EVIDENCE</b>", styles['Heading2']))
    Story.append(Spacer(1, 6))
    Story.append(Paragraph(f"URL:<br/><a href='{reporte.url_evidencia}' color='blue'>{reporte.url_evidencia}</a>", normal_style))
    Story.append(Spacer(1, 10))
    
    compressed_img_io = fetch_and_compress_image(reporte)
    if compressed_img_io:
        try:
            rl_img = RLImage(compressed_img_io)
            rl_img.drawHeight = 250
            rl_img.drawWidth = 250
            rl_img.hAlign = 'LEFT'
            Story.append(rl_img)
        except Exception:
            Story.append(Paragraph("<i>(La imagen no pudo ser incrustada)</i>", normal_style))
    else:
        Story.append(Paragraph("<i>(Sin imagen disponible capturada por el bot)</i>", normal_style))
    
    doc.build(Story)
    buffer.seek(0)
    return buffer

def generate_investigation_docx(db: Session, reporte, perfil):
    doc = Document()
    
    doc.add_heading(f"INTELLIGENCE REPORT — {reporte.pais.upper()}", 0)
    doc.add_paragraph(f"Ref: {reporte.referencia_interna} | ID producto: {reporte.id}")
    
    doc.add_heading('COMPANY OVERVIEW', level=1)
    
    username = reporte.titular_cuenta
    arroba = ""
    username_clean = username.replace("@", "") if username else "Desconocido"
    owner_display = f"@{username_clean}"

    doc.add_heading('FASE 1 DE LA INVESTIGACIÓN (RECOLECCIÓN)', level=2)
    
    table1 = doc.add_table(rows=1, cols=2)
    table1.style = 'Table Grid'
    hdr_cells1 = table1.rows[0].cells
    hdr_cells1[0].text = "Campo (Fase 1: Recolección)"
    hdr_cells1[1].text = "Valor detectado"
    
    fields_fase1 = [
        ("País", reporte.pais.upper()),
        ("Referencia interna", reporte.referencia_interna),
        ("Owner (Titular)", owner_display),
        ("Social Networks / Plataforma", reporte.plataforma),
        ("Website / URL evidencia", reporte.url_evidencia),
        ("Fecha de escaneo", reporte.fecha_captura.strftime("%Y-%m-%d %H:%M:%S") if reporte.fecha_captura else "N/D"),
        ("Proposed fees (Precio)", f"USD {reporte.precio_publico}" if reporte.precio_publico else "USD —"),
        ("Catalogue (Producto)", reporte.producto_catalogo),
        ("Marca afectada", reporte.marca_afectada),
        ("Sales (Hallazgos)", f"Según la investigación, se detectó oferta de productos con indicio de infracción a la marca {reporte.marca_afectada}. Origen: Automático."),
        ("Type of Seller", "Por determinar"),
        ("Estimated quantity", "Por determinar"),
        ("Address Evidence", reporte.pais)
    ]
    
    for k, v in fields_fase1:
        row_cells = table1.add_row().cells
        row_cells[0].text = str(k)
        row_cells[1].text = str(v)
        
    doc.add_paragraph()
    doc.add_heading('FASE 2 DE RELEVAMIENTO (IDENTIDAD Y TRAZABILIDAD)', level=2)
    
    table2 = doc.add_table(rows=1, cols=2)
    table2.style = 'Table Grid'
    hdr_cells2 = table2.rows[0].cells
    hdr_cells2[0].text = "Campo (Fase 2: Relevamiento de Identidad)"
    hdr_cells2[1].text = "Valor detectado"
    
    fields_fase2 = [
        ("DNI / Cédula", getattr(perfil, 'dni_cedula', None) or "—"),
        ("CUIT / CUIL / RUC", getattr(perfil, 'cuit_cuil_ruc', None) or "—"),
        ("Banco", getattr(perfil, 'banco', None) or "—"),
        ("Alias bancario / MercadoPago", getattr(perfil, 'alias_bancario', None) or "—"),
        ("CVU / CBU", getattr(perfil, 'cvu_cbu', None) or "—"),
        ("Address (Dirección exacta)", getattr(perfil, 'direccion_exacta', None) or "—"),
        ("Telephone Number (Privado)", getattr(perfil, 'telefono_privado', None) or "—"),
        ("Email", getattr(perfil, 'email', None) or "—"),
        ("Precio final cotizado", getattr(perfil, 'precio_final_cotizado', None) or "—")
    ]
    
    for k, v in fields_fase2:
        row_cells = table2.add_row().cells
        row_cells[0].text = str(k)
        row_cells[1].text = str(v)
        
    doc.add_heading('ADDRESS / PHOTO EVIDENCE', level=1)
    doc.add_paragraph(f"URL: {reporte.url_evidencia}")
    
    compressed_img_io = fetch_and_compress_image(reporte)
    if compressed_img_io:
        try:
            doc.add_picture(compressed_img_io, width=Inches(3.5))
        except Exception:
            doc.add_paragraph("(La imagen no pudo ser incrustada)")
    else:
        doc.add_paragraph("(Sin imagen disponible capturada por el bot)")
    
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def generate_reportes_list_pdf(reportes):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    
    Story = []
    Story.append(Paragraph('Listado de Reportes Net Holmes', styles['Heading1']))
    Story.append(Spacer(1, 12))
    
    from reportlab.platypus import Table, TableStyle
    from reportlab.lib import colors
    
    data = [['ID', 'Marca', 'Plataforma', 'Vendedor', 'Estado']]
    for r in reportes:
        data.append([str(r.id), r.marca_afectada, r.plataforma, r.titular_cuenta[:20], r.estado])
        
    t = Table(data, colWidths=[40, 100, 100, 150, 100])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.grey),
        ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,0), 10),
        ('BOTTOMPADDING', (0,0), (-1,0), 12),
        ('BACKGROUND', (0,1), (-1,-1), colors.beige),
        ('GRID', (0,0), (-1,-1), 1, colors.black)
    ]))
    
    Story.append(t)
    doc.build(Story)
    buffer.seek(0)
    return buffer

def generate_reportes_list_docx(reportes):
    doc = Document()
    doc.add_heading('Listado de Reportes Net Holmes', 0)
    
    table = doc.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'ID'
    hdr_cells[1].text = 'Marca'
    hdr_cells[2].text = 'Plataforma'
    hdr_cells[3].text = 'Vendedor'
    hdr_cells[4].text = 'Estado'
    
    for r in reportes:
        row_cells = table.add_row().cells
        row_cells[0].text = str(r.id)
        row_cells[1].text = r.marca_afectada
        row_cells[2].text = r.plataforma
        row_cells[3].text = r.titular_cuenta
        row_cells[4].text = r.estado
        
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
