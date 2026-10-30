"""Negocio, almacenamiento y generación documental de Seizure Reports."""

from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
import json
import re
import uuid
import zipfile

from fastapi import HTTPException, UploadFile
from PIL import Image
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Image as RLImage, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import or_
from sqlalchemy.orm import Session
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from models.brand import Brand
from models.seizure_report import (
    SeizureAuditEvent, SeizureDocumentVersion, SeizureFile, SeizureHistory, SeizureProduct,
    SeizureReport, SeizureSequence, SeizureTarget,
)
from services.procedure_reference_service import assign_reference, brand_code as official_brand_code, normalize_report_type


STORAGE = Path(__file__).resolve().parent.parent / "assets" / "seizure_reports"
STORAGE.mkdir(parents=True, exist_ok=True)
TEMPLATE_DIR = Path(__file__).resolve().parent.parent / "assets" / "seizure_templates"
NAVY, CYAN = colors.HexColor("#172f53"), colors.HexColor("#10c5c9")
DISCLAIMER = (
    "The information contained in this report has been obtained through lawful and publicly available sources, "
    "in compliance with applicable regulations of the Argentine Republic, including Article 43 of the National "
    "Constitution, Law No. 25,326 on Personal Data Protection, Law No. 27,401 on Corporate Criminal Liability and "
    "related compliance standards. This report is strictly limited to the information, images, records and documents "
    "expressly included herein and is provided solely for internal compliance, intellectual property protection, "
    "risk assessment and enforcement purposes. Underlying records and working papers remain confidential."
)
BRAND_CODES = {
    "puma":"PU", "new era":"NE", "lacoste":"LC", "tommy hilfiger":"TH", "calvin klein":"CK",
    "under armour":"UA", "mlb":"MLB", "nba":"NBA", "wella company":"WC", "polo":"PRL",
    "ralph lauren":"PRL", "among us":"AU", "universal studios":"US", "lego":"LE",
    "victoria's secret":"VP", "ray ban":"RB", "chanel":"CH", "alo":"AL", "nike":"NI",
}
STATUSES = {
    "BORRADOR", "PENDIENTE_REVISION", "OBSERVADO", "APROBADO_INTERNAMENTE", "REGISTRADO",
    "PENDIENTE_APROBACION_CLIENTE", "APROBADO_CLIENTE", "RECHAZADO_CLIENTE", "ARCHIVADO", "ENVIADO_FACTURACION",
}
TRANSITIONS = {
    "BORRADOR": {"PENDIENTE_REVISION", "ARCHIVADO"},
    "PENDIENTE_REVISION": {"OBSERVADO", "APROBADO_INTERNAMENTE"},
    "OBSERVADO": {"BORRADOR", "PENDIENTE_REVISION"},
    "APROBADO_INTERNAMENTE": {"REGISTRADO"},
    "REGISTRADO": {"PENDIENTE_APROBACION_CLIENTE", "ARCHIVADO"},
    "PENDIENTE_APROBACION_CLIENTE": {"APROBADO_CLIENTE", "RECHAZADO_CLIENTE"},
    "RECHAZADO_CLIENTE": {"OBSERVADO", "ARCHIVADO"},
    "APROBADO_CLIENTE": {"ENVIADO_FACTURACION", "ARCHIVADO"},
    "ENVIADO_FACTURACION": {"ARCHIVADO"}, "ARCHIVADO": set(),
}


def brand_code(name: str | None) -> str:
    return official_brand_code(name)


def get_report(db: Session, report_id: int) -> SeizureReport:
    report = db.query(SeizureReport).filter_by(id=report_id).first()
    if not report:
        raise HTTPException(404, "Informe de allanamiento no encontrado")
    return report


def list_reports(db, *, search=None, status=None, brand_id=None, country=None, skip=0, limit=20):
    query = db.query(SeizureReport)
    if search:
        value = f"%{search.strip()}%"
        query = query.filter(or_(SeizureReport.internal_reference.ilike(value), SeizureReport.client_case_number.ilike(value)))
    if status:
        query = query.filter(SeizureReport.status == status)
    if brand_id:
        query = query.filter(SeizureReport.brand_id == brand_id)
    if country:
        query = query.filter(SeizureReport.country == country)
    total = query.count()
    items = query.order_by(SeizureReport.updated_at.desc()).offset(skip).limit(limit).all()
    return {"items": items, "total": total}


def audit_event(db: Session, report_id: int, user_id: int, action: str, *,
                entity_type: str = "report", entity_id: int | None = None, details=None):
    row = SeizureAuditEvent(
        report_id=report_id, user_id=user_id, action=action,
        entity_type=entity_type, entity_id=entity_id,
        details=json.dumps(details, ensure_ascii=False, default=str) if details is not None else None,
    )
    db.add(row)
    return row


def apply_payload(db: Session, report: SeizureReport, data, user_id: int):
    is_new = report.id is None
    data.report_type = normalize_report_type(data.report_type)
    scalar = data.model_dump(exclude={"targets", "products"})
    if data.brand_id:
        brand = db.get(Brand, data.brand_id)
        if not brand:
            raise HTTPException(422, "La marca seleccionada no existe")
        scalar["brand_name"] = brand.nombre
        scalar["brand_code"] = (data.brand_code or brand_code(brand.nombre)).upper()
    for key, value in scalar.items():
        setattr(report, key, value)
    report.targets[:] = [SeizureTarget(**target.model_dump()) for target in data.targets]
    report.products[:] = [SeizureProduct(**product.model_dump()) for product in data.products]
    report.seized_goods_total = sum(product.quantity for product in data.products)
    if is_new:
        db.add(report)
        db.flush()
        db.add(SeizureHistory(report_id=report.id, to_status="BORRADOR", comment="Borrador creado", user_id=user_id))
    audit_event(db, report.id, user_id, "REPORT_CREATED" if is_new else "REPORT_UPDATED", details={
        "targets": len(data.targets), "products": len(data.products),
        "seized_goods_total": report.seized_goods_total,
    })
    db.commit(); db.refresh(report)
    return report


def transition(db, report, target, comment, user):
    target = target.upper()
    if target not in STATUSES or target not in TRANSITIONS.get(report.status, set()):
        raise HTTPException(409, f"Transición no permitida: {report.status} → {target}")
    administrative = {
        "OBSERVADO", "APROBADO_INTERNAMENTE", "REGISTRADO", "PENDIENTE_APROBACION_CLIENTE",
        "APROBADO_CLIENTE", "RECHAZADO_CLIENTE", "ENVIADO_FACTURACION",
    }
    if target in administrative and user.rol not in {"admin", "sup_admin"}:
        raise HTTPException(403, "Esta transición requiere permisos de administrador")
    if target in {"PENDIENTE_REVISION", "APROBADO_INTERNAMENTE"}:
        missing = []
        photos = [row for row in report.files if row.kind == "photo"]
        documents = [row for row in report.files if row.kind == "document"]
        if not photos: missing.append("fotografías de la operación")
        if not documents: missing.append("acta o documentación policial")
        if not report.targets: missing.append("domicilio u objetivo")
        if not report.products or report.seized_goods_total <= 0: missing.append("cantidades incautadas")
        if missing:
            raise HTTPException(409, "Información pendiente: " + ", ".join(missing))
    if target == "APROBADO_INTERNAMENTE":
        pending_products = [row for row in report.products if "pendiente" in (row.brand_name or "").casefold()]
        distinct_clients = {row.client_name or row.brand_name for row in report.products if row.client_name or row.brand_name}
        unclassified_photos = [row for row in report.files if row.kind == "photo" and not row.brand_name]
        if pending_products:
            raise HTTPException(409, "Debe distribuir las cantidades pendientes por cliente o marca antes de aprobar")
        if len(distinct_clients) > 1 and unclassified_photos:
            raise HTTPException(409, "Debe clasificar todas las fotografías por cliente o marca antes de aprobar")
    old = report.status
    report.status = target
    db.add(SeizureHistory(report_id=report.id, from_status=old, to_status=target, comment=comment, user_id=user.id))
    audit_event(db, report.id, user.id, "STATUS_CHANGED", details={"from": old, "to": target, "comment": comment})
    if target == "APROBADO_INTERNAMENTE" and not report.internal_reference:
        report.internal_reference, report.brand_code = assign_reference(
            db, report.report_type, report.brand_name or report.client_name,
            brand_code_value=report.brand_code,
        )
    db.commit(); db.refresh(report)
    return report


async def store_upload(db, report, upload: UploadFile, kind: str, user_id: int, position: int = 0):
    allowed = {"photo": {"image/jpeg", "image/png", "image/webp"}, "document": {"application/pdf", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}}
    if upload.content_type not in allowed[kind]:
        raise HTTPException(415, f"Formato no permitido: {upload.content_type}")
    content = await upload.read()
    if not content or len(content) > 20 * 1024 * 1024:
        raise HTTPException(413, "El archivo debe pesar entre 1 byte y 20 MB")
    suffix = Path(upload.filename or "file").suffix.lower()
    stored = f"{report.id}_{uuid.uuid4().hex}{suffix}"
    target = STORAGE / stored
    target.write_bytes(content)
    if kind == "photo":
        try:
            with Image.open(target) as image: image.verify()
        except Exception:
            target.unlink(missing_ok=True); raise HTTPException(415, "La imagen no es válida")
    elif upload.content_type == "application/pdf":
        try:
            PdfReader(BytesIO(content))
        except Exception:
            target.unlink(missing_ok=True); raise HTTPException(415, "El PDF no es válido")
    elif not zipfile.is_zipfile(BytesIO(content)):
        target.unlink(missing_ok=True); raise HTTPException(415, "El documento Word no es válido")
    row = SeizureFile(report_id=report.id, kind=kind, original_name=Path(upload.filename or "archivo").name,
                      stored_name=stored, mime_type=upload.content_type, size_bytes=len(content), position=position,
                      uploaded_by=user_id, document_type="ACTA_POLICIAL" if kind == "document" else None)
    db.add(row); db.flush()
    audit_event(db, report.id, user_id, "FILE_UPLOADED", entity_type="file", entity_id=row.id,
                details={"kind": kind, "name": row.original_name, "size_bytes": row.size_bytes})
    db.commit(); db.refresh(row)
    return row


def extract_police_pdf(path_or_bytes) -> dict:
    reader = PdfReader(path_or_bytes)
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    clean_text = text.replace("\r", "\n")
    def one(pattern, flags=re.I):
        match = re.search(pattern, clean_text, flags)
        return " ".join(match.group(1).split()) if match else None
    def section(label, following):
        boundary = "|".join(following)
        return one(rf"(?:^|\n)\s*{label}\s*:\s*(.*?)(?=\n\s*(?:{boundary})\s*:|$)", re.I | re.S)

    date_text = one(r"(?:^|\n)\s*FECHA(?:\s+DE\s+LA\s+OPERACI[ÓO]N)?\s*:\s*([^\n]+)")
    iso_date = None
    numeric_date = re.search(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", date_text or "")
    if numeric_date:
        iso_date = f"{numeric_date.group(3)}-{int(numeric_date.group(2)):02d}-{int(numeric_date.group(1)):02d}"

    places_text = section(r"LUGAR(?:ES)?", [r"IMPUTADOS?", r"PERSONAL(?:\s+AFECTADO)?", r"MAGISTRADO(?:\s+INTERVENTOR)?", r"ORGANISMOS", r"RECURSOS", r"SECUESTRO"])
    places = []
    if places_text:
        numbered = re.findall(r"(?:^|\n)\s*\d+[).]\s*(.+?)(?=\n\s*\d+[).]|$)", places_text, re.S)
        places = [" ".join(value.split()) for value in numbered] if numbered else [" ".join(places_text.split())]

    infringers = section(r"IMPUTADOS?", [r"MAGISTRADO(?:\s+INTERVENTOR)?", r"RECURSOS", r"SECUESTRO", r"RESE[ÑN]A"])
    court = section(r"MAGISTRADO(?:\s+INTERVENTOR)?", [r"CAUSA", r"RECURSOS", r"SECUESTRO", r"RESE[ÑN]A"])
    seizure = section(r"SECUESTRO", [r"RESE[ÑN]A", r"OBSERVACIONES", r"ESTADO"])
    summary = section(r"RESE[ÑN]A", [r"ESTADO", r"FIRMA"])
    amount = None
    if seizure:
        amount_match = re.search(r"\(\s*([\d.]+)\s*\)", seizure)
        if amount_match:
            amount = int(amount_match.group(1).replace(".", ""))
    known_brands = sorted({name.title() for name in BRAND_CODES} | {"Nike", "Jordan", "Adidas", "Levis", "Converse", "Gucci", "Dior", "Versace", "Diesel", "Boss"}, key=len, reverse=True)
    identified_brands = [name for name in known_brands if re.search(rf"(?<!\w){re.escape(name)}(?!\w)", seizure or "", re.I)]
    products = []
    if amount:
        products.append({
            "brand_name": "Pendiente de clasificación por cliente",
            "category": "Vestimenta y calzado",
            "description": f"Total informado por el acta. Marcas mencionadas: {', '.join(identified_brands) or 'sin detalle'}",
            "quantity": amount, "unit": "unidades",
            "notes": "Cantidad total confirmada; distribución por cliente pendiente de validación humana",
        })

    heading_lines = [" ".join(line.split()) for line in clean_text.splitlines()[:8] if line.strip()]
    agency = " / ".join(line for line in heading_lines if re.fullmatch(r"(?:DCC|DIC|DIVISI[ÓO]N|DIRECCI[ÓO]N).{0,80}", line, re.I)) or None
    return {
        "operation_date_text": date_text, "operation_date": iso_date,
        "case_origin": section(r"ASUNTO", [r"FECHA", r"HORA(?:\s+DE\s+INICIO)?", r"LUGARES?"]),
        "police_agency": agency or one(r"POLICE REPORT.*?\n([^\n]+)", re.I | re.S),
        "police_unit": section(r"RECURSOS", [r"SECUESTRO", r"RESE[ÑN]A"]),
        "court": court or one(r"JUZGADO\s*:\s*([^\n]+)"),
        "court_case_number": one(r"CAUSA\s*:\s*([^\n]+)") or one(r"\b(OS\s*\d+/\d{4})\b"),
        "targets": places, "infringers": infringers, "seized_goods_total": amount,
        "identified_brands": identified_brands, "products": products,
        "operation_summary": summary, "media_impact": "N/A" if re.search(r"negativo\s+medios", clean_text, re.I) else None,
        "raw_text": clean_text[:20000],
    }


def _photo_dimensions(path, max_w, max_h):
    with Image.open(path) as image: w, h = image.size
    scale = min(max_w / w, max_h / h)
    return w * scale, h * scale


def _report_values(report):
    targets = "; ".join(filter(None, [f"{t.establishment_name + ': ' if t.establishment_name else ''}{t.address}" for t in report.targets])) or "N/A"
    breakdown = {}
    for product in getattr(report, "products", []):
        key = product.client_name or product.brand_name or "Pending classification"
        breakdown[key] = breakdown.get(key, 0) + (product.quantity or 0)
    breakdown_text = "; ".join(f"{name}: {quantity}" for name, quantity in breakdown.items()) or "N/A"
    return [
        ("TARGET", report.brand_name or report.client_name or "N/A"), ("DATE", report.operation_date.strftime("%d/%m/%Y") if report.operation_date else "N/A"),
        ("ORIGIN OF THE CASE", report.case_origin or "N/A"), ("SEARCH AND WARRANT TARGETS", targets),
        ("AMOUNT OF SEIZED GOODS", str(report.seized_goods_total)), ("SEIZED GOODS BY CLIENT / BRAND", breakdown_text),
        ("INFRINGERS", "; ".join(t.investigated_party for t in report.targets if t.investigated_party) or "N/A"),
        ("LAW ENFORCEMENT AGENCY", " - ".join(filter(None, [report.police_agency, report.police_unit])) or "N/A"),
        ("JUDGE", " - ".join(filter(None, [report.court, report.judge])) or "N/A"), ("MEDIA", report.media_impact or "N/A"),
        ("POLITICS", report.political_impact or "N/A"), ("PROPOSED FEES", f"{report.proposed_fees or 'N/A'} {report.currency or ''}".strip()),
        ("INTERNAL REFERENCE", report.internal_reference or "N/A"),
    ]


def _photo_caption(photo, index):
    parts = [getattr(photo, "caption", None), getattr(photo, "brand_name", None), getattr(photo, "establishment", None)]
    return " · ".join(value for value in parts if value) or f"Evidence {index + 1}"


def _report_sections(report):
    values = dict(_report_values(report))
    return [
        ("TARGET", [(key, values[key]) for key in ("DATE", "ORIGIN OF THE CASE", "SEARCH AND WARRANT TARGETS", "AMOUNT OF SEIZED GOODS", "SEIZED GOODS BY CLIENT / BRAND", "INFRINGERS")]),
        ("AUTHORITIES", [(key, values[key]) for key in ("LAW ENFORCEMENT AGENCY", "JUDGE")]),
        ("EXTERNAL IMPACT", [(key, values[key]) for key in ("MEDIA", "POLITICS")]),
        ("INTERNAL INFORMATION", [(key, values[key]) for key in ("PROPOSED FEES", "INTERNAL REFERENCE")]),
    ]


def _shade_cell(cell, fill):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _set_cell_text(cell, text, *, color="000000", bold=False, size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _add_corporate_table(doc, title, rows):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Cm(5.2), Cm(11.5)
    heading = table.rows[0].cells
    heading[0].merge(heading[1])
    _set_cell_text(heading[0], title, color="1E4459", bold=True, size=11)
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].width, cells[1].width = Cm(5.2), Cm(11.5)
        _shade_cell(cells[0], "1E4459")
        _set_cell_text(cells[0], label, color="FFFFFF", bold=True)
        _set_cell_text(cells[1], value, size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def generate_pdf(report, user_id):
    output = STORAGE / f"report_{report.id}_{uuid.uuid4().hex}.pdf"
    styles = getSampleStyleSheet(); styles.add(ParagraphStyle(name="Section", parent=styles["Heading1"], textColor=NAVY, fontSize=18, spaceAfter=14))
    styles.add(ParagraphStyle(name="WhiteBody", parent=styles["BodyText"], textColor=colors.white))
    styles.add(ParagraphStyle(name="WhiteLabel", parent=styles["BodyText"], textColor=colors.white, fontName="Helvetica-Bold"))
    disclaimer = Table([
        [Paragraph("COMPLIANCE DISCLAIMER", styles["WhiteLabel"])],
        [Paragraph(DISCLAIMER, styles["WhiteBody"])],
    ], colWidths=[16.7*cm])
    disclaimer.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#C00000")),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
        ("BOX", (0, 0), (-1, -1), .5, colors.HexColor("#8A0000")),
        ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story = [Spacer(1, 1), PageBreak(), Paragraph(f"{report.country.upper()} – SEIZURE REPORT", styles["Section"]), disclaimer, Spacer(1, .5*cm)]
    rows, section_rows = [], []
    for title, values in _report_sections(report):
        section_rows.append(len(rows)); rows.append([Paragraph(f"<b>{title}</b>", styles["BodyText"]), ""])
        rows.extend([Paragraph(key, styles["WhiteLabel"]), Paragraph(str(value), styles["BodyText"])] for key, value in values)
    table = Table(rows, colWidths=[5.2*cm, 11.5*cm], repeatRows=0)
    commands = [("GRID",(0,0),(-1,-1),.35,colors.HexColor("#9aa7b7")),("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),6),("RIGHTPADDING",(0,0),(-1,-1),6),("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6)]
    for index in range(len(rows)):
        if index in section_rows:
            commands += [("SPAN",(0,index),(1,index)),("TEXTCOLOR",(0,index),(1,index),NAVY),("BACKGROUND",(0,index),(1,index),colors.white)]
        else:
            commands += [("BACKGROUND",(0,index),(0,index),NAVY),("TEXTCOLOR",(0,index),(0,index),colors.white)]
    table.setStyle(TableStyle(commands)); story += [table]
    def pdf_section_heading(title):
        heading = Table([[Paragraph(title, styles["WhiteLabel"])]], colWidths=[16.7*cm])
        heading.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), NAVY), ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
            ("BOX", (0, 0), (-1, -1), .5, NAVY), ("LEFTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ]))
        return heading
    photos = [f for f in report.files if f.kind == "photo" and f.include_in_report]
    if photos: story += [PageBreak(), pdf_section_heading("PICTURES OF THE OPERATION"), Spacer(1, .35*cm)]
    for index, photo in enumerate(photos):
        path = STORAGE/photo.stored_name; w,h = _photo_dimensions(path, 16.5*cm, 10.5*cm)
        story += [RLImage(str(path), width=w, height=h), Paragraph(_photo_caption(photo, index), styles["BodyText"]), Spacer(1,.4*cm)]
        if index % 2 == 1 and index < len(photos)-1: story.append(PageBreak())
    documents = [f for f in report.files if f.kind == "document" and f.include_in_report and f.mime_type == "application/pdf"]
    if documents:
        story += [PageBreak(), pdf_section_heading("POLICE REPORT: CONFIDENTIAL"), Spacer(1, .4*cm),
                  Paragraph("The original police documentation is attached in full on the following pages.", styles["BodyText"])]
    def page_decoration(canvas, document):
        canvas.saveState()
        cover = TEMPLATE_DIR / "cover.png"
        if document.page == 1 and cover.exists():
            canvas.drawImage(str(cover), 0, 0, width=A4[0], height=A4[1], preserveAspectRatio=True, mask="auto")
            canvas.restoreState()
            return
        logo = TEMPLATE_DIR / "cba-logo.png"
        if logo.exists():
            canvas.drawImage(str(logo), 1.7*cm, A4[1]-1.2*cm, width=3.6*cm, height=.8*cm, preserveAspectRatio=True, mask="auto")
        canvas.setFillColor(NAVY); canvas.setFont("Helvetica", 7)
        canvas.drawRightString(A4[0]-1.7*cm, A4[1]-.85*cm, "CORPORATE INTELLIGENCE IN ACTION")
        canvas.setStrokeColor(CYAN); canvas.setLineWidth(1)
        canvas.line(1.7*cm, 1.25*cm, A4[0]-1.7*cm, 1.25*cm)
        canvas.setFillColor(NAVY); canvas.setFont("Helvetica-Bold", 8)
        canvas.drawString(1.7*cm, .8*cm, "CONFIDENTIAL REPORT")
        canvas.drawRightString(A4[0]-1.7*cm, .8*cm, f"PAGE {document.page}")
        canvas.restoreState()
    SimpleDocTemplate(str(output), pagesize=A4, rightMargin=1.7*cm,leftMargin=1.7*cm,topMargin=1.6*cm,bottomMargin=1.6*cm, title="Seizure Report").build(story, onFirstPage=page_decoration, onLaterPages=page_decoration)
    if documents:
        merged = PdfWriter();
        for page in PdfReader(str(output)).pages: merged.add_page(page)
        for attachment in documents:
            for page in PdfReader(str(STORAGE/attachment.stored_name)).pages: merged.add_page(page)
        with output.open("wb") as stream: merged.write(stream)
    return output


def generate_docx(report, user_id):
    output = STORAGE / f"report_{report.id}_{uuid.uuid4().hex}.docx"; doc = Document()
    section=doc.sections[0]; section.top_margin=section.bottom_margin=Cm(1.7); section.left_margin=section.right_margin=Cm(1.7)
    header=section.header.paragraphs[0]
    logo=TEMPLATE_DIR/"cba-logo.png"
    if logo.exists(): header.add_run().add_picture(str(logo),width=Cm(3.6))
    header.add_run("\tCORPORATE INTELLIGENCE IN ACTION"); header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer=section.footer.paragraphs[0]; footer.text="CONFIDENTIAL REPORT"; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    cover=TEMPLATE_DIR/"cover.png"
    if cover.exists():
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(cover),width=Cm(17.5),height=Cm(24.75))
    else:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("INTELLIGENCE REPORT\n"); r.bold=True; r.font.size=Pt(34); r.font.color.rgb=RGBColor(23,47,83); r=p.add_run("SEIZURE REPORT · CONFIDENTIAL"); r.bold=True; r.font.size=Pt(17); r.font.color.rgb=RGBColor(16,197,201)
    doc.add_page_break()
    title = doc.add_paragraph()
    title.style = "Intense Quote"
    title_run = title.add_run(f"{report.country.upper()} – SEIZURE REPORT")
    title_run.bold = True; title_run.font.name = "Arial"; title_run.font.size = Pt(18); title_run.font.color.rgb = RGBColor.from_string("1E4459")
    disclaimer = doc.add_table(rows=2, cols=1); disclaimer.style = "Table Grid"
    _shade_cell(disclaimer.cell(0, 0), "C00000"); _shade_cell(disclaimer.cell(1, 0), "C00000")
    _set_cell_text(disclaimer.cell(0, 0), "COMPLIANCE DISCLAIMER", color="FFFFFF", bold=True, size=9)
    _set_cell_text(disclaimer.cell(1, 0), DISCLAIMER, color="FFFFFF", size=7)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    for section_title, rows in _report_sections(report):
        _add_corporate_table(doc, section_title, rows)
    photos=[f for f in report.files if f.kind=="photo" and f.include_in_report]
    if photos:
        doc.add_page_break()
        photo_heading = doc.add_table(rows=1, cols=1); photo_heading.style = "Table Grid"
        _shade_cell(photo_heading.cell(0, 0), "1E4459")
        _set_cell_text(photo_heading.cell(0, 0), "PICTURES OF THE OPERATION", color="FFFFFF", bold=True, size=12)
    for index,photo in enumerate(photos):
        path=STORAGE/photo.stored_name; w,h=_photo_dimensions(path,15.8*cm,10*cm); doc.add_picture(str(path),width=Cm(w/cm),height=Cm(h/cm)); doc.add_paragraph(_photo_caption(photo,index))
        if index%2==1 and index<len(photos)-1: doc.add_page_break()
    documents=[f for f in report.files if f.kind=="document" and f.include_in_report and f.mime_type=="application/pdf"]
    if documents:
        doc.add_page_break()
        police_heading = doc.add_table(rows=1, cols=1); police_heading.style = "Table Grid"
        _shade_cell(police_heading.cell(0, 0), "1E4459")
        _set_cell_text(police_heading.cell(0, 0), "POLICE REPORT: CONFIDENTIAL", color="FFFFFF", bold=True, size=12)
    for attachment in documents:
        try: text="\n".join(page.extract_text() or "" for page in PdfReader(str(STORAGE/attachment.stored_name)).pages); doc.add_heading(attachment.original_name,2); doc.add_paragraph(text)
        except Exception: doc.add_paragraph(f"Attached document: {attachment.original_name}")
    doc.save(output); return output


def register_version(db, report, path, fmt, user_id):
    version = (db.query(SeizureDocumentVersion).filter_by(report_id=report.id, format=fmt).count() + 1)
    row=SeizureDocumentVersion(report_id=report.id,version=version,format=fmt,stored_name=path.name,generated_by=user_id)
    db.add(row); db.flush()
    audit_event(db, report.id, user_id, "DOCUMENT_GENERATED", entity_type="document_version", entity_id=row.id,
                details={"format": fmt, "version": version, "file": path.name})
    db.commit()
    return row
